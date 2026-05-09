from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.config import settings


class ExcelReportGenerationError(Exception):
    """Raised when Excel report generation fails."""


CURRENCY_FORMAT = '#,##0.00;[Red](#,##0.00);"-"'
PERCENT_FORMAT = '0.00%;[Red](0.00%);"-"'
NUMBER_FORMAT = '#,##0.00;[Red](#,##0.00);"-"'
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=14, bold=True, color="1F4E78")
THIN_BORDER = Border(bottom=Side(style="thin", color="D9E2F3"))


def generate_excel_report(
    extracted_metrics: dict[str, Any],
    calculated_ratios: dict[str, Any],
    ai_insights: list[dict[str, Any]] | dict[str, Any],
    output_dir: str | Path | None = None,
    filename: str | None = None,
) -> str:
    metrics = _extract_list(extracted_metrics, "metrics")
    ratios = _extract_list(calculated_ratios, "ratios")
    insights = ai_insights.get("insights") if isinstance(ai_insights, dict) else ai_insights
    if not isinstance(insights, list):
        raise ExcelReportGenerationError("ai_insights must be a list or contain an insights list.")

    destination_dir = Path(output_dir or settings.export_dir)
    destination_dir.mkdir(parents=True, exist_ok=True)
    destination = destination_dir / (filename or f"financial_report_{uuid4().hex}.xlsx")
    if destination.suffix.lower() != ".xlsx":
        destination = destination.with_suffix(".xlsx")

    workbook = Workbook()
    workbook.remove(workbook.active)

    _write_executive_summary(workbook, metrics, ratios, insights)
    _write_extracted_financials(workbook, metrics)
    _write_ratios(workbook, ratios)
    _write_ai_insights(workbook, insights)
    _write_source_references(workbook, metrics, insights)

    try:
        workbook.save(destination)
    except OSError as exc:
        raise ExcelReportGenerationError(f"Failed to save Excel report: {destination}") from exc

    return str(destination)


def _extract_list(payload: dict[str, Any], key: str) -> list[dict[str, Any]]:
    values = payload.get(key)
    if not isinstance(values, list):
        raise ExcelReportGenerationError(f"Input must include a {key} list.")
    if not all(isinstance(value, dict) for value in values):
        raise ExcelReportGenerationError(f"Each {key} item must be a JSON object.")
    return values


def _write_executive_summary(
    workbook: Workbook,
    metrics: list[dict[str, Any]],
    ratios: list[dict[str, Any]],
    insights: list[dict[str, Any]],
) -> None:
    sheet = workbook.create_sheet("Executive Summary")
    sheet["A1"] = "Executive Summary"
    sheet["A1"].font = TITLE_FONT
    sheet["A3"] = "Generated At"
    sheet["B3"] = datetime.now(UTC).replace(microsecond=0).isoformat()
    sheet["A4"] = "Extracted Metrics"
    sheet["B4"] = len(metrics)
    sheet["A5"] = "Calculated Ratios"
    sheet["B5"] = len(ratios)
    sheet["A6"] = "AI Insights"
    sheet["B6"] = len(insights)

    _write_rows(
        sheet,
        start_row=8,
        headers=["Metric", "Current Year", "Previous Year", "Unit", "Source Page"],
        rows=[
            [
                metric.get("metric_name"),
                metric.get("current_year_value"),
                metric.get("previous_year_value"),
                metric.get("unit"),
                metric.get("source_page"),
            ]
            for metric in metrics[:8]
        ],
    )

    ratio_start = max(sheet.max_row + 3, 18)
    _write_rows(
        sheet,
        start_row=ratio_start,
        headers=["Ratio", "Value", "Unit", "Explanation"],
        rows=[
            [
                ratio.get("ratio_name"),
                _percentage_decimal(ratio.get("value")) if ratio.get("unit") == "%" else ratio.get("value"),
                ratio.get("unit"),
                ratio.get("explanation"),
            ]
            for ratio in ratios[:8]
        ],
    )

    _style_sheet(sheet, freeze_cell="A8")
    _apply_currency_format(sheet, columns=[2, 3], start_row=9)
    _apply_percentage_format(sheet, columns=[2], start_row=ratio_start + 1)


def _write_extracted_financials(workbook: Workbook, metrics: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet("Extracted Financials")
    _write_rows(
        sheet,
        start_row=1,
        headers=[
            "Metric Name",
            "Current Year Value",
            "Previous Year Value",
            "Unit",
            "Source Page",
            "Confidence Score",
            "Source Text",
        ],
        rows=[
            [
                metric.get("metric_name"),
                metric.get("current_year_value"),
                metric.get("previous_year_value"),
                metric.get("unit"),
                metric.get("source_page"),
                metric.get("confidence_score"),
                metric.get("source_text"),
            ]
            for metric in metrics
        ],
    )
    _style_sheet(sheet)
    _apply_currency_format(sheet, columns=[2, 3], start_row=2)
    _apply_percentage_format(sheet, columns=[6], start_row=2)


def _write_ratios(workbook: Workbook, ratios: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet("Ratios")
    _write_rows(
        sheet,
        start_row=1,
        headers=[
            "Ratio Name",
            "Value",
            "Unit",
            "Formula",
            "Missing Inputs",
            "Explanation",
        ],
        rows=[
            [
                ratio.get("ratio_name"),
                _percentage_decimal(ratio.get("value")) if ratio.get("unit") == "%" else ratio.get("value"),
                ratio.get("unit"),
                ratio.get("formula"),
                _join_values(ratio.get("missing_inputs")),
                ratio.get("explanation"),
            ]
            for ratio in ratios
        ],
    )
    _style_sheet(sheet)
    _apply_percentage_format(sheet, columns=[2], start_row=2)


def _write_ai_insights(workbook: Workbook, insights: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet("AI Insights")
    _write_rows(
        sheet,
        start_row=1,
        headers=[
            "Title",
            "Type",
            "Summary",
            "Supporting Metrics",
            "Source Pages",
            "Confidence Score",
        ],
        rows=[
            [
                insight.get("title"),
                insight.get("type"),
                insight.get("summary"),
                _join_values(insight.get("supporting_metrics")),
                _join_values(insight.get("source_pages")),
                insight.get("confidence_score"),
            ]
            for insight in insights
        ],
    )
    _style_sheet(sheet)
    _apply_percentage_format(sheet, columns=[6], start_row=2)


def _write_source_references(
    workbook: Workbook,
    metrics: list[dict[str, Any]],
    insights: list[dict[str, Any]],
) -> None:
    sheet = workbook.create_sheet("Source References")
    metric_rows = [
        [
            metric.get("source_page"),
            "metric",
            metric.get("metric_name"),
            metric.get("source_text"),
        ]
        for metric in metrics
    ]
    insight_rows = [
        [
            _join_values(insight.get("source_pages")),
            "insight",
            insight.get("title"),
            insight.get("summary"),
        ]
        for insight in insights
    ]
    _write_rows(
        sheet,
        start_row=1,
        headers=["Source Page", "Reference Type", "Reference Name", "Source Text / Summary"],
        rows=metric_rows + insight_rows,
    )
    _style_sheet(sheet)


def _write_rows(
    sheet: Worksheet,
    start_row: int,
    headers: list[str],
    rows: list[list[Any]],
) -> None:
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=start_row, column=column_index, value=header)
        _style_header_cell(cell)

    for row_offset, row in enumerate(rows, start=start_row + 1):
        for column_index, value in enumerate(row, start=1):
            sheet.cell(row=row_offset, column=column_index, value=value)


def _style_sheet(sheet: Worksheet, freeze_cell: str = "A2") -> None:
    sheet.freeze_panes = freeze_cell
    for row in sheet.iter_rows():
        for cell in row:
            if cell.fill.fill_type == "solid" and cell.font.bold:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER

    _auto_width(sheet)


def _style_header_cell(cell: Any) -> None:
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _auto_width(sheet: Worksheet) -> None:
    for column_cells in sheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 60)


def _apply_currency_format(sheet: Worksheet, columns: list[int], start_row: int) -> None:
    for column in columns:
        for row in range(start_row, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=column)
            if isinstance(cell.value, int | float):
                cell.number_format = CURRENCY_FORMAT


def _apply_percentage_format(sheet: Worksheet, columns: list[int], start_row: int) -> None:
    for column in columns:
        for row in range(start_row, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=column)
            if isinstance(cell.value, int | float):
                cell.number_format = PERCENT_FORMAT


def _percentage_decimal(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, int | float):
        return value / 100
    return value


def _join_values(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return ", ".join(str(item) for item in value)
    return str(value)


def generate_multi_year_excel(
    comparison_payload: dict[str, Any],
    output_dir: str | Path | None = None,
    filename: str | None = None,
) -> str:
    """Write multi-document comparison workbook (Overview, metrics, ratios, AI summary)."""

    destination_dir = Path(output_dir or settings.export_dir)
    destination_dir.mkdir(parents=True, exist_ok=True)
    destination = destination_dir / (filename or f"comparison_{uuid4().hex}.xlsx")
    if destination.suffix.lower() != ".xlsx":
        destination = destination.with_suffix(".xlsx")

    workbook = Workbook()
    workbook.remove(workbook.active)

    comparison_hash = str(comparison_payload.get("comparison_hash") or "")
    documents = comparison_payload.get("documents") or []
    metric_series = comparison_payload.get("metric_series") or []
    ratio_series = comparison_payload.get("ratio_series") or []
    trend_summary = comparison_payload.get("trend_summary") or {}

    _write_comparison_overview(workbook, comparison_hash, documents)
    _write_comparison_metric_trends(workbook, metric_series, documents)
    _write_comparison_ratio_trends(workbook, ratio_series, documents)
    _write_comparison_trend_summary(workbook, trend_summary)

    try:
        workbook.save(destination)
    except OSError as exc:
        raise ExcelReportGenerationError(f"Failed to save multi-year Excel report: {destination}") from exc

    return str(destination)


def _write_comparison_overview(
    workbook: Workbook,
    comparison_hash: str,
    documents: list[Any],
) -> None:
    sheet = workbook.create_sheet("Overview")
    sheet["A1"] = "Multi-Year Comparison"
    sheet["A1"].font = TITLE_FONT
    sheet["A3"] = "Comparison ID"
    sheet["B3"] = comparison_hash
    sheet["A4"] = "Generated At"
    sheet["B4"] = datetime.now(UTC).replace(microsecond=0).isoformat()

    headers = ["Document ID", "File Name", "Company", "Report Year"]
    rows = []
    for d in documents:
        if isinstance(d, dict):
            rows.append(
                [
                    d.get("document_id"),
                    d.get("file_name"),
                    d.get("company_name"),
                    d.get("report_year"),
                ]
            )

    _write_rows(sheet, start_row=6, headers=headers, rows=rows)
    _style_sheet(sheet, freeze_cell="A7")


def _write_comparison_metric_trends(
    workbook: Workbook,
    metric_series: list[Any],
    documents_meta: list[Any],
) -> None:
    sheet = workbook.create_sheet("Metric Trends")

    doc_cols = _sorted_document_columns(documents_meta)

    headers = ["Metric", "Unit"] + [
        f"Year {y} (#{did})" if y is not None else f"Doc #{did}"
        for did, y in doc_cols
    ]
    headers += ["YoY % (latest vs prior)"]

    rows: list[list[Any]] = []
    for series in metric_series:
        if not isinstance(series, dict):
            continue
        name = series.get("metric_name")
        unit = series.get("unit")
        points = series.get("points") or []
        point_map = _points_by_doc_year(points)

        row: list[Any] = [name, unit]
        for did, yr in doc_cols:
            key = (did, yr)
            p = point_map.get(key)
            row.append(p.get("value") if p else None)

        row.append(_latest_yoy_percent(points))
        rows.append(row)

    _write_rows(sheet, start_row=1, headers=headers, rows=rows)
    _style_sheet(sheet)


def _write_comparison_ratio_trends(
    workbook: Workbook,
    ratio_series: list[Any],
    documents_meta: list[Any],
) -> None:
    sheet = workbook.create_sheet("Ratio Trends")
    doc_cols = _sorted_document_columns(documents_meta)

    headers = ["Ratio", "Unit"] + [
        f"Year {y} (#{did})" if y is not None else f"Doc #{did}"
        for did, y in doc_cols
    ]
    headers += ["YoY % (latest vs prior)"]

    rows: list[list[Any]] = []
    for series in ratio_series:
        if not isinstance(series, dict):
            continue
        name = series.get("ratio_name")
        unit = series.get("unit")
        points = series.get("points") or []
        point_map = _points_by_doc_year(points)

        row: list[Any] = [name, unit]
        for did, yr in doc_cols:
            key = (did, yr)
            p = point_map.get(key)
            row.append(p.get("value") if p else None)

        row.append(_latest_yoy_percent(points))
        rows.append(row)

    _write_rows(sheet, start_row=1, headers=headers, rows=rows)
    _style_sheet(sheet)


def _write_comparison_trend_summary(workbook: Workbook, trend_summary: Any) -> None:
    sheet = workbook.create_sheet("Trend Summary")
    if not isinstance(trend_summary, dict):
        trend_summary = {}

    sheet["A1"] = "AI Trend Summary"
    sheet["A1"].font = TITLE_FONT
    sheet["A3"] = "Headline"
    sheet["B3"] = trend_summary.get("headline")
    sheet["A4"] = "Confidence"
    sheet["B4"] = trend_summary.get("confidence")

    bullets = trend_summary.get("bullets") or []
    if isinstance(bullets, list):
        sheet["A6"] = "Bullets"
        sheet["A6"].font = TITLE_FONT
        for idx, bullet in enumerate(bullets, start=7):
            sheet.cell(row=idx, column=1, value=f"{idx - 6}.")
            sheet.cell(row=idx, column=2, value=bullet)

    _style_sheet(sheet, freeze_cell="A7")


def _sorted_document_columns(documents_meta: list[Any]) -> list[tuple[int, int | None]]:
    enriched: list[tuple[int, int | None]] = []
    for d in documents_meta:
        if not isinstance(d, dict):
            continue
        did = d.get("document_id")
        yr = d.get("report_year")
        if isinstance(did, int):
            enriched.append((did, yr if isinstance(yr, int) else None))

    enriched.sort(key=lambda item: (item[1] if item[1] is not None else -1, item[0]))
    return enriched


def _points_by_doc_year(points: Any) -> dict[tuple[int, int | None], dict[str, Any]]:
    out: dict[tuple[int, int | None], dict[str, Any]] = {}
    if not isinstance(points, list):
        return out
    for p in points:
        if not isinstance(p, dict):
            continue
        did = p.get("document_id")
        yr = p.get("report_year")
        if not isinstance(did, int):
            continue
        yrint = yr if isinstance(yr, int) else None
        out[(did, yrint)] = p
    return out


def _latest_yoy_percent(points: Any) -> float | None:
    if not isinstance(points, list):
        return None
    sorted_pts = sorted(
        [p for p in points if isinstance(p, dict)],
        key=lambda p: (
            p.get("report_year") if isinstance(p.get("report_year"), int) else -1,
            p.get("document_id") or 0,
        ),
    )
    if not sorted_pts:
        return None
    return sorted_pts[-1].get("yoy_percent")
